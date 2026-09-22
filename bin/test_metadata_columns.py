"""A web edit must not delete the columns a lab added to sample_metadata.xlsx.

`sample_metadata.xlsx` is offered for Download / Replace so "a lab can manage
it locally", and a lab that does adds columns of its own. The mirror used to be
REBUILT from the seven canonical fields on every web save, which deleted them.
It is now edited in place.

Run directly:  <conda>/bin/python bin/test_metadata_columns.py
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl  # noqa: E402

import metadata as meta  # noqa: E402

FAILED = 0


def check(label, cond):
    global FAILED
    if cond:
        print(f"  OK  {label}")
    else:
        FAILED += 1
        print(f"  FAIL {label}")


def grid(path):
    ws = openpyxl.load_workbook(path).active
    return {(c.row, c.column): c.value
            for row in ws.iter_rows() for c in row if c.value is not None}


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="irma_meta_"))

    print("[a first write still produces our own labeled workbook]")
    fresh = tmp / "sample_metadata.xlsx"
    meta.write_xlsx({"S1": {"sample": "S1", "host": "chicken"}}, fresh)
    hdr = [c.value for c in openpyxl.load_workbook(fresh).active[1]]
    check("the header is the canonical field list", hdr == meta.FIELDS)
    check("a second write round-trips it unchanged in shape",
          (meta.write_xlsx({"S1": {"sample": "S1", "host": "duck"}}, fresh)
           and [c.value for c in openpyxl.load_workbook(fresh).active[1]] == meta.FIELDS))
    check("...and the edited field did change", grid(fresh)[(2, 3)] == "duck")

    print("[a lab's own workbook keeps its own columns]")
    lab = tmp / "lab.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    for r, row in enumerate([
        ["sample", "host", "submitter", "strain", "review notes"],
        ["S1", "chicken", "Vivek", "A/chicken/PA/S1/2026", "resequence requested"],
        ["S2", "turkey", "Tod", "A/turkey/PA/S2/2026", "clean"],
        [None, None, None, None, "reviewed 2026-09-01"],
    ], start=1):
        for c, v in enumerate(row, start=1):
            if v is not None:
                ws.cell(row=r, column=c, value=v)
    wb.save(lab)

    meta.write_xlsx({
        "S1": {"sample": "S1", "host": "chicken", "state": "PA",
               "collection_year": "2026", "subtype": "H5N1"},
        "S3": {"sample": "S3", "host": "mallard", "state": "MD"},
    }, lab)
    g = grid(lab)
    hdr = [c.value for c in openpyxl.load_workbook(lab).active[1]]

    check("the lab's columns are all still in the header",
          hdr[:5] == ["sample", "host", "submitter", "strain", "review notes"])
    check("their cells are untouched",
          g[(2, 3)] == "Vivek" and g[(2, 5)] == "resequence requested"
          and g[(3, 3)] == "Tod" and g[(3, 5)] == "clean")
    check("a free-text 'strain' column is NOT overwritten with the subtype",
          g[(2, 4)] == "A/chicken/PA/S1/2026")
    check("the subtype went to a column of its own",
          g[(2, hdr.index("subtype") + 1)] == "H5N1")
    check("the lab's existing host column was updated in place, not duplicated",
          hdr.count("host") == 1 and g[(2, 2)] == "chicken")
    check("the trailing note row was not written over",
          g[(4, 5)] == "reviewed 2026-09-01" and (4, 1) not in g)
    check("a sample not yet in the sheet was appended below it",
          g[(5, 1)] == "S3" and g[(5, 2)] == "mallard")
    check("a row for a sample the web UI did not send is left alone",
          g[(3, 1)] == "S2" and g[(3, 2)] == "turkey")

    print("[a workbook with no sample column is left alone, not replaced]")
    opaque = tmp / "opaque.xlsx"
    wb = openpyxl.Workbook()
    wb.active.cell(row=1, column=1, value="notes only")
    wb.active.cell(row=2, column=1, value="nothing matchable here")
    wb.save(opaque)
    before = opaque.read_bytes()
    try:
        meta.write_xlsx({"S1": {"sample": "S1"}}, opaque)
        refused = False
    except ValueError:
        refused = True
    check("the write is refused rather than guessed at", refused)
    check("...and the file is byte-for-byte unchanged", opaque.read_bytes() == before)

    print("PASS" if not FAILED else f"{FAILED} FAILURE(S)")
    return 1 if FAILED else 0


if __name__ == "__main__":
    raise SystemExit(main())
